"""Citeste PI-uri de la furnizori (xlsx/xls/csv): antet pe 1-2 randuri, celule unite, dimensiuni pe 3 coloane,
G.W./N.W. pe carton, plus date generale (furnizor, nr. PI, data, moneda, transport, incoterm)."""
import re

import numpy as np
import pandas as pd

SYNONYMS = {
    "sku": ["item number", "item no", "item", "model", "sku", "article", "product code", "part no", "ref", "code", "cod"],
    "description": ["description", "product name", "item name", "goods", "descriere", "name", "product"],
    "qty": ["total pcs", "total qty", "quantity", "qty", "order qty", "pcs", "units", "cantitate"],
    "unit_price": ["unit price", "u/p", "price/pc", "unit cost", "fob", "price", "pret"],
    "amount": ["total amount", "amount", "total", "value", "valoare"],
    "gw": ["g w /ctns", "g w /ctn", "gross weight", "gross wt", "g w", "gw"],
    "nw": ["n w /ctns", "n w /ctn", "net weight", "net wt", "n w", "nw"],
    "weight": ["weight", "greutate"],
    "ppc": ["pcs/ctns", "pcs/ctn", "pcs per carton", "qty/ctn", "pcs/carton", "units/ctn", "packing", "buc/carton", "per carton"],
    "cartons": ["cartons", "ctns", "ctn", "no of ctn", "nr cutii"],
    "size": ["carton size", "ctn size", "carton dimension", "carton measurement", "measurement"],
    "hs": ["hs code", "hs", "customs code", "tariff", "nc"],
}
FIELD_LABELS = {"sku": "SKU / Item no.", "description": "Descriere", "qty": "Cantitate (total buc)", "unit_price": "Preț unitar",
                "amount": "Valoare totală (opțional)", "ppc": "Buc / carton", "cartons": "Nr. cartoane (opțional)",
                "size": "Dimensiuni carton (L, l, H)", "gw": "Greutate brută (G.W.)", "nw": "Greutate netă (N.W.)", "weight": "Greutate generică", "hs": "Cod NC/HS"}
WEIGHT_CHOICES = {"Brută per carton (G.W.)": ("gw", "carton"), "Netă per carton (N.W.)": ("nw", "carton"),
                  "Per bucată": ("weight", "piece"), "Total linie": ("weight", "total")}
CARTON_ONLY = ("gw", "nw", "size", "ppc")


def _norm(x):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9/ ]+", " ", str(x).lower())).strip()


def _score(cell, syn):
    c, s = _norm(cell), _norm(syn)
    if not c or str(cell) == "nan": return 0
    if c == s: return 30 + len(s)
    if c.startswith(s) or c.endswith(s): return 20 + len(s)
    if len(s) >= 3 and s in c: return 10 + len(s)
    return 0


def _isna(v): return v is None or (isinstance(v, float) and np.isnan(v)) or (not isinstance(v, str) and pd.isna(v))


# ------------------------------------------------------------------ citire fisier
def read_raw(file):
    """{foaie: DataFrame brut}. Pentru xlsx pastreaza si celulele unite in df.attrs['merges'] (0-based)."""
    name = getattr(file, "name", str(file)).lower()
    if name.endswith(".csv"):
        return {"csv": pd.read_csv(file, header=None, dtype=object, sep=None, engine="python")}
    if name.endswith(".xls"):
        xl = pd.ExcelFile(file)
        return {sh: xl.parse(sh, header=None, dtype=object) for sh in xl.sheet_names}
    from openpyxl import load_workbook
    wb = load_workbook(file, data_only=True)
    out = {}
    for ws in wb:
        df = pd.DataFrame([list(r) for r in ws.iter_rows(values_only=True)], dtype=object)
        df.attrs["merges"] = [(m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges]
        out[ws.title] = df
    return out


def _cell(raw, r, j):
    v = raw.iat[r, j]
    if _isna(v):
        for r1, c1, r2, c2 in raw.attrs.get("merges", []):
            if r1 <= r <= r2 and c1 <= j <= c2:
                return raw.iat[r1, c1]
    return v


def detect_header(raw, max_rows=40):
    """(rand_antet, nr_randuri_antet). Al doilea rand de antet = sub-titluri text sub randul principal."""
    best, best_n = 0, -1
    for i in range(min(max_rows, len(raw))):
        n = sum(1 for v in raw.iloc[i] if not _isna(v) and any(_score(v, s) for ss in SYNONYMS.values() for s in ss))
        if n > best_n: best, best_n = i, n
    nxt = raw.iloc[best + 1] if best + 1 < len(raw) else []
    txt = [v for v in nxt if not _isna(v) and isinstance(v, str)]
    nums = [v for v in nxt if not _isna(v) and not isinstance(v, str)]
    return best, (2 if len(txt) >= 2 and not nums else 1)


def with_header(raw, hdr, n=1):
    """Tabel cu antet: eticheta = rand inferior (daca exista) altfel cel superior; celulele unite se extind; etichetele repetate primesc #2, #3."""
    labels, seen = [], {}
    for j in range(raw.shape[1]):
        parts = [_cell(raw, r, j) for r in range(hdr, hdr + n)]
        parts = [str(p).strip() for p in parts if not _isna(p) and str(p).strip()]
        lab = parts[-1] if parts else f"col{j}"
        seen[lab] = seen.get(lab, 0) + 1
        labels.append(lab if seen[lab] == 1 else f"{lab}#{seen[lab]}")
    df = raw.iloc[hdr + n:].copy()
    df.columns = labels
    return df.dropna(how="all").reset_index(drop=True)


def auto_map(columns):
    """camp -> nume coloana. Coloanele 'Total ...' nu pot fi greutate/dimensiune; o coloana serveste un singur camp."""
    cand = []
    for f, syns in SYNONYMS.items():
        for idx, c in enumerate(columns):
            if f in CARTON_ONLY and _norm(c).startswith("total"): continue
            sc = max((_score(re.sub(r"#\d+$", "", c), s) for s in syns), default=0)
            if sc: cand.append((sc, -idx, f, c))
    mapping, used = {}, set()
    for sc, _, f, c in sorted(cand, reverse=True):
        if f not in mapping and c not in used:
            mapping[f] = c; used.add(c)
    return mapping


# ------------------------------------------------------------------ date generale
def detect_meta(raw):
    """Furnizor, nr. PI, data, moneda, transport, incoterm - cautate prin eticheta din apropiere sau text."""
    meta = {}
    cells = [(r, j, raw.iat[r, j]) for r in range(min(len(raw), 60)) for j in range(raw.shape[1]) if not _isna(raw.iat[r, j])]

    def right_of(r, j):
        for k in range(j + 1, raw.shape[1]):
            if not _isna(raw.iat[r, k]) and str(raw.iat[r, k]).strip(): return raw.iat[r, k]
    for r, j, v in cells:
        n = _norm(v)
        if n == "date" and "date" not in meta:
            d = pd.to_datetime(right_of(r, j), errors="coerce")
            if pd.notna(d): meta["date"] = d
        elif n in ("invoice", "invoice no", "invoice number", "pi no", "pi number", "proforma invoice no") and "pi_no" not in meta:
            x = right_of(r, j)
            if x is not None: meta["pi_no"] = str(x).strip()
        elif n == "currency" and "currency" not in meta:
            x = str(right_of(r, j) or "").upper()
            if x in ("USD", "EUR", "RON"): meta["currency"] = x
        elif n.startswith("freight type") and "transport_type" not in meta:
            x = _norm(right_of(r, j) or "")
            meta["transport_type"] = ("Aerian" if "air" in x else "Maritim LCL" if ("sea" in x or "ocean" in x) else
                                      "Rutier" if ("road" in x or "truck" in x) else "Curier" if ("express" in x or "courier" in x) else "Aerian")
        m = re.search(r"incoterms?\W*(EXW|FOB|CIF|CFR|DDP|DAP|FCA|CIP|CPT|DPU|FAS)", str(v), re.I)
        if m and "incoterm" not in meta: meta["incoterm"] = m.group(1).upper()
    for r, j, v in cells:
        if r < 5 and isinstance(v, str) and len(v) > 5 and not re.search(r"invoice|proforma|pro forma", v, re.I):
            meta["supplier"] = v.strip(); break
    if "currency" not in meta:
        head = " ".join(str(v) for r, j, v in cells if r < 30 and re.search(r"price|amount", str(v), re.I)).lower()
        meta["currency"] = "EUR" if ("eur" in head or "€" in head) else "USD" if ("us$" in head or "usd" in head or "$" in head) else None
    return {k: v for k, v in meta.items() if v}


# ------------------------------------------------------------------ normalizare
_DIM = re.compile(r"(\d+(?:[.,]\d+)?)\s*[x*×X]\s*(\d+(?:[.,]\d+)?)\s*[x*×X]\s*(\d+(?:[.,]\d+)?)")


def parse_dims(v):
    m = _DIM.search(str(v))
    return tuple(float(g.replace(",", ".")) for g in m.groups()) if m else (np.nan,) * 3


def _num(s): return pd.to_numeric(s.astype(str).str.replace(",", "").str.extract(r"(-?\d+\.?\d*)")[0], errors="coerce")


def normalize(df, mapping, weight_field="gw", weight_basis="carton", dim_unit="cm", assume_one_carton=False):
    """Linii standard. weight_basis: 'carton' | 'piece' | 'total'. Daca lipseste buc/carton si nu e dedus, ramane gol (se completeaza in editor)."""
    blank = pd.Series([np.nan] * len(df), index=df.index)
    g = lambda f: df[mapping[f]] if mapping.get(f) in df.columns else blank
    txt = lambda f: g(f).fillna("").astype(str).str.strip().replace("nan", "")
    out = pd.DataFrame({"sku": txt("sku"), "description": txt("description"), "hs_code": txt("hs"), "qty": _num(g("qty")),
                        "unit_price": _num(g("unit_price"))})
    amount, cartons, ppc = _num(g("amount")), _num(g("cartons")), _num(g("ppc"))
    out["unit_price"] = out["unit_price"].fillna(amount / out["qty"].replace(0, np.nan))
    ppc = ppc.fillna(out["qty"] / cartons.replace(0, np.nan))
    if assume_one_carton: ppc = ppc.fillna(out["qty"])
    out["pcs_per_carton"] = ppc
    w = _num(g(weight_field))
    w = w / ppc.replace(0, np.nan) if weight_basis == "carton" else (w / out["qty"].replace(0, np.nan) if weight_basis == "total" else w)
    out["unit_weight_kg"] = w
    k = {"cm": 1.0, "mm": 0.1, "m": 100.0}[dim_unit]
    base = re.sub(r"#\d+$", "", mapping.get("size", "")) if mapping.get("size") else None
    grp = [c for c in df.columns if base and re.sub(r"#\d+$", "", c) == base]
    if len(grp) >= 3:
        for c, col in zip(("carton_l", "carton_w", "carton_h"), grp[:3]): out[c] = _num(df[col]) * k
    else:
        dims = g("size").map(parse_dims)
        for i, c in enumerate(("carton_l", "carton_w", "carton_h")): out[c] = dims.map(lambda t: t[i]) * k
    out = out[(out["qty"] > 0) & (out["sku"] != "") & ~out["sku"].str.lower().str.contains("total")]
    return out.reset_index(drop=True)
